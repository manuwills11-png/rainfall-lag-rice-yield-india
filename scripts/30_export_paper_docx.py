"""Convert docs/research_paper.md to docs/research_paper.docx (headings, tables, figures, bullets, code, bold/italic)."""
import re
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt

ROOT = Path(__file__).resolve().parent.parent
SRC = ROOT / "docs" / "research_paper.md"


def add_inline(par, text):
    for tok in re.split(r"(\*\*[^*]+\*\*|\*[^*\s][^*]*\*|`[^`]+`)", text):
        if not tok:
            continue
        if tok.startswith("**") and tok.endswith("**"):
            par.add_run(tok[2:-2]).bold = True
        elif tok.startswith("`") and tok.endswith("`"):
            r = par.add_run(tok[1:-1]); r.font.name = "Consolas"; r.font.size = Pt(9)
        elif tok.startswith("*") and tok.endswith("*") and len(tok) > 2:
            par.add_run(tok[1:-1]).italic = True
        else:
            par.add_run(tok)


def main():
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"; doc.styles["Normal"].font.size = Pt(10.5)
    lines = SRC.read_text(encoding="utf8").splitlines()
    i = 0
    while i < len(lines):
        ln = lines[i]
        if ln.startswith("```"):
            j = i + 1; buf = []
            while not lines[j].startswith("```"):
                buf.append(lines[j]); j += 1
            p = doc.add_paragraph(); r = p.add_run("\n".join(buf)); r.font.name = "Consolas"; r.font.size = Pt(8.5)
            i = j + 1; continue
        if ln.startswith("|") and i + 1 < len(lines) and re.match(r"^\|[\s:\-|]+\|$", lines[i + 1]):
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                if not re.match(r"^\|[\s:\-|]+\|$", lines[i]):
                    rows.append([c.strip() for c in lines[i].strip().strip("|").split("|")])
                i += 1
            t = doc.add_table(rows=len(rows), cols=len(rows[0])); t.style = "Light Grid Accent 1"
            for r_, row in enumerate(rows):
                for c_, cell in enumerate(row):
                    par = t.cell(r_, c_).paragraphs[0]; add_inline(par, cell)
                    for run in par.runs:
                        run.font.size = Pt(8.5); run.bold = run.bold or r_ == 0
            doc.add_paragraph(); continue
        m = re.match(r"^(#{1,4})\s+(.*)", ln)
        if m:
            doc.add_heading(re.sub(r"[*`]", "", m.group(2)), level=min(len(m.group(1)), 3) - (0 if len(m.group(1)) > 1 else 0)); i += 1; continue
        img = re.search(r"!\[[^\]]*\]\(([^)]+)\)", ln)
        if img:
            path = (SRC.parent / img.group(1)).resolve()
            pre = ln[: img.start()].strip()
            if pre:
                add_inline(doc.add_paragraph(), pre)
            if path.exists():
                doc.add_picture(str(path), width=Inches(6.3))
            post = ln[img.end():].strip()
            if post:
                add_inline(doc.add_paragraph(), post)
            i += 1; continue
        if re.match(r"^\s*[*-]\s+", ln):
            add_inline(doc.add_paragraph(style="List Bullet"), re.sub(r"^\s*[*-]\s+", "", ln)); i += 1; continue
        if ln.strip() in ("---", ""):
            i += 1; continue
        add_inline(doc.add_paragraph(), ln); i += 1
    out = ROOT / "docs" / "research_paper.docx"; doc.save(out); print("wrote", out)


if __name__ == "__main__":
    main()
